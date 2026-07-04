from pathlib import Path

from openpyxl import Workbook
from docx import Document

from app.parsers import parse_docx_file, parse_text_file, parse_xlsx_file


def test_parse_text_file_reads_utf8(tmp_path: Path):
    path = tmp_path / "note.md"
    path.write_text("# 标题\n专题验收助手", encoding="utf-8")

    assert "专题验收助手" in parse_text_file(path)


def test_parse_docx_file_reads_paragraphs(tmp_path: Path):
    path = tmp_path / "note.docx"
    document = Document()
    document.add_paragraph("需求管理")
    document.save(str(path))

    assert "需求管理" in parse_docx_file(path)


def test_parse_xlsx_file_reads_sheets(tmp_path: Path):
    path = tmp_path / "plan.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "进度"
    sheet.append(["任务", "状态"])
    sheet.append(["专题验收", "进行中"])
    workbook.save(str(path))

    result = parse_xlsx_file(path)
    assert "[Sheet: 进度]" in result
    assert "专题验收" in result
